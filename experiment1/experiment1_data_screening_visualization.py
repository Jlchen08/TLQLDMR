"""
Experiment 1: Extreme-Weather Scenario Screening & Visualization

Goal:
  1) For each wind farm, compare normal vs extreme samples (distribution shift).
  2) (Optional) Show extreme-event time-series snippets for a selected farm.

Notes:
  - This script intentionally avoids pandas/sklearn to run in environments where those
    wheels may be incompatible with the current NumPy version.
  - Excel reading uses openpyxl; KDE uses scipy.
  - All outputs are saved under TL-QLDMR/experiment1/.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import numpy as np
import matplotlib.pyplot as plt
import openpyxl
from scipy.stats import gaussian_kde


# -----------------------------
# Data loading (xlsx -> numpy)
# -----------------------------


def _to_float(x) -> float:
    if x is None:
        return float("nan")
    if isinstance(x, (int, float, np.floating)):
        return float(x)
    # openpyxl may return datetime/string for time column; caller should not float it
    try:
        return float(x)
    except Exception:
        return float("nan")


def _to_time_str(x) -> str:
    if x is None:
        return ""
    if isinstance(x, datetime):
        return x.isoformat(sep=" ")
    if isinstance(x, str):
        # Try to normalize common "YYYY-mm-dd HH:MM:SS" strings
        s = x.strip()
        try:
            dt = datetime.fromisoformat(s)
            return dt.isoformat(sep=" ")
        except Exception:
            return s
    return str(x)


def _find_column(headers: List[Optional[str]], keywords: Iterable[str]) -> Optional[int]:
    kws = [k.lower() for k in keywords]
    for j, h in enumerate(headers):
        if not h:
            continue
        hs = str(h).lower()
        if all(k in hs for k in kws):
            return j
    return None


def _parse_capacity_mw(stem: str) -> Optional[int]:
    # Example: "Wind farm site 2 (Nominal capacity-200MW)"
    import re

    m = re.search(r"capacity-(\d+)mw", stem.lower())
    if not m:
        return None
    return int(m.group(1))


@dataclass(frozen=True)
class FarmInfo:
    path: Path
    name: str
    capacity_mw: int


def list_wind_farms(data_dir: Path) -> List[FarmInfo]:
    farms: List[FarmInfo] = []
    for p in sorted(data_dir.glob("*.xlsx")):
        cap = _parse_capacity_mw(p.stem)
        if cap is None:
            continue
        farms.append(FarmInfo(path=p, name=p.stem, capacity_mw=cap))
    farms.sort(key=lambda f: f.capacity_mw, reverse=True)
    return farms


def load_farm_timeseries_xlsx(
    xlsx_path: Path,
    feature_set: str,
) -> Tuple[Dict[str, np.ndarray], Dict[str, str]]:
    """
    Returns:
      series: dict with keys: time, power, wind_hub, temperature, ... (depending on feature_set)
      col_names: mapping of logical key -> original excel header (for provenance)
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

        col_time = _find_column(headers, ["time", "year-month-day"])
        col_power = _find_column(headers, ["power", "mw"])
        col_wind_hub = _find_column(headers, ["wind speed", "wheel hub", "m/s"])
        col_wind_50m = _find_column(headers, ["wind speed", "50 meters", "m/s"])
        col_wind_30m = _find_column(headers, ["wind speed", "30 meters", "m/s"])
        col_wind_10m = _find_column(headers, ["wind speed", "10 meters", "m/s"])
        col_temp = _find_column(headers, ["air temperature"])
        col_pressure = _find_column(headers, ["atmosphere", "hpa"])
        col_humidity = _find_column(headers, ["relative humidity"])

        if col_power is None:
            raise ValueError(f"Cannot find Power column in {xlsx_path.name}")
        if col_wind_hub is None:
            raise ValueError(f"Cannot find hub-height wind speed column in {xlsx_path.name}")

        # Decide which feature columns to load (time series; no sliding window here).
        feature_set = feature_set.lower().strip()
        if feature_set not in {"full", "wind_only", "simple"}:
            raise ValueError("feature_set must be one of: full, wind_only, simple")

        needed: Dict[str, Optional[int]] = {
            "time": col_time,
            "power": col_power,
            "wind_hub": col_wind_hub,
        }
        if feature_set == "full":
            needed.update(
                {
                    "wind_50m": col_wind_50m,
                    "wind_30m": col_wind_30m,
                    "wind_10m": col_wind_10m,
                    "temperature": col_temp,
                    "pressure": col_pressure,
                    "humidity": col_humidity,
                }
            )
        elif feature_set == "wind_only":
            needed.update(
                {
                    "wind_50m": col_wind_50m,
                    "wind_30m": col_wind_30m,
                    "wind_10m": col_wind_10m,
                }
            )
        else:  # simple
            needed.update({"temperature": col_temp})

        # Filter out missing columns (keep for metadata but don't read values).
        present = {k: v for k, v in needed.items() if v is not None}

        # Preallocate lists; openpyxl read-only iteration is row-wise.
        out: Dict[str, List] = {k: [] for k in present.keys()}
        for row in ws.iter_rows(min_row=2, values_only=True):
            for k, j in present.items():
                v = row[j] if j < len(row) else None
                if k == "time":
                    out[k].append(_to_time_str(v))
                else:
                    out[k].append(_to_float(v))

        series = {k: np.asarray(v) for k, v in out.items()}
        col_names = {k: (headers[j] if j is not None else "") for k, j in needed.items()}
        return series, col_names
    finally:
        wb.close()


# -----------------------------
# Extreme event identification
# -----------------------------


@dataclass(frozen=True)
class ExtremeMasks:
    extreme: np.ndarray  # union mask for target domain
    ramp: np.ndarray
    cutout: np.ndarray
    stat: np.ndarray
    temp: np.ndarray
    ramp_threshold_mw: float
    q95_wind: float
    q95_power: float


def identify_extreme_events(
    power_mw: np.ndarray,
    wind_hub_ms: np.ndarray,
    temperature_c: Optional[np.ndarray],
    nominal_capacity_mw: float,
) -> ExtremeMasks:
    power = np.asarray(power_mw, dtype=float)
    wind = np.asarray(wind_hub_ms, dtype=float)
    temp = None if temperature_c is None else np.asarray(temperature_c, dtype=float)

    q95_wind = float(np.nanpercentile(wind, 95))
    q95_power = float(np.nanpercentile(power, 95))

    stat = (wind > q95_wind) | (power > q95_power)

    ramp_threshold = 0.05 * float(nominal_capacity_mw)
    power_diff = np.abs(np.diff(power, prepend=power[0]))
    ramp = power_diff > ramp_threshold

    cutout = wind > 25.0

    if temp is None:
        temp_mask = np.zeros_like(power, dtype=bool)
    else:
        temp_mask = (temp < -10.0) | (temp > 35.0)

    extreme = stat | ramp | cutout | temp_mask
    return ExtremeMasks(
        extreme=extreme,
        ramp=ramp,
        cutout=cutout,
        stat=stat,
        temp=temp_mask,
        ramp_threshold_mw=float(ramp_threshold),
        q95_wind=q95_wind,
        q95_power=q95_power,
    )


# -----------------------------
# Visualizations
# -----------------------------


def _select_non_overlapping_segments(
    event_idx: np.ndarray,
    n: int,
    pre: int,
    post: int,
    total_len: int,
) -> List[Tuple[int, int, int]]:
    """
    Returns segments as (event_i, start, end) with end exclusive.
    """
    picked: List[Tuple[int, int, int]] = []
    last_end = -10**9
    for i in event_idx:
        i = int(i)
        start = max(0, i - pre)
        end = min(total_len, i + post + 1)
        if start < last_end:
            continue
        if end - start < (pre + post) // 2:
            continue
        picked.append((i, start, end))
        last_end = end
        if len(picked) >= n:
            break
    return picked


def plot_timeseries_segment(
    out_path: Path,
    time_str: np.ndarray,
    wind_hub: np.ndarray,
    power: np.ndarray,
    event_i: int,
    start: int,
    end: int,
    title: str,
) -> None:
    x = np.arange(start, end)
    t0 = time_str[start] if len(time_str) else ""
    t1 = time_str[end - 1] if len(time_str) else ""

    fig, ax1 = plt.subplots(figsize=(12, 5))
    ax2 = ax1.twinx()

    ax1.plot(x, wind_hub[start:end], color="tab:blue", lw=1.5, label="Wind speed (hub, m/s)")
    ax2.plot(x, power[start:end], color="tab:orange", lw=1.5, label="Power (MW)")

    ax1.axvline(event_i, color="red", ls="--", lw=1.0, alpha=0.8)
    ax1.set_xlabel("Index (15-min steps)")
    ax1.set_ylabel("Wind speed (m/s)")
    ax2.set_ylabel("Power (MW)")

    subtitle = f"[{t0} ~ {t1}]  event_index={event_i}"
    ax1.set_title(f"{title}\n{subtitle}")

    # Combined legend
    lines = ax1.get_lines() + ax2.get_lines()
    labels = [l.get_label() for l in lines]
    ax1.legend(lines, labels, loc="upper right")
    ax1.grid(True, alpha=0.25)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.tight_layout()
    fig.savefig(out_path, dpi=300)
    plt.close(fig)


def plot_kde_shift(
    out_path: Path,
    source: np.ndarray,
    target: np.ndarray,
    xlabel: str,
    title: str,
    bw_method: Optional[float] = None,
) -> None:
    s = np.asarray(source, dtype=float)
    t = np.asarray(target, dtype=float)
    s = s[np.isfinite(s)]
    t = t[np.isfinite(t)]

    out_path.parent.mkdir(parents=True, exist_ok=True)

    fig, ax = plt.subplots(figsize=(8.5, 4.5))
    if len(s) < 2 or len(t) < 2:
        ax.text(
            0.5,
            0.5,
            f"Not enough samples for KDE\nsource={len(s)}, target={len(t)}",
            ha="center",
            va="center",
        )
        ax.set_axis_off()
        fig.savefig(out_path, dpi=300)
        plt.close(fig)
        return

    lo = float(min(np.min(s), np.min(t)))
    hi = float(max(np.max(s), np.max(t)))
    if math.isclose(lo, hi):
        hi = lo + 1.0
    pad = 0.05 * (hi - lo)
    xs = np.linspace(lo - pad, hi + pad, 400)

    kde_s = gaussian_kde(s, bw_method=bw_method)
    kde_t = gaussian_kde(t, bw_method=bw_method)

    ax.plot(xs, kde_s(xs), lw=2.0, label=f"Source (normal) n={len(s)}")
    ax.plot(xs, kde_t(xs), lw=2.0, label=f"Target (extreme) n={len(t)}")
    ax.set_title(title)
    ax.set_xlabel(xlabel)
    ax.set_ylabel("Density")
    ax.grid(True, alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(out_path, dpi=300)
    plt.close(fig)


def _maybe_subsample(x: np.ndarray, max_samples: int, rng: np.random.Generator) -> np.ndarray:
    if len(x) <= max_samples:
        return x
    idx = rng.choice(len(x), size=max_samples, replace=False)
    return x[idx]


def _get_palette(n: int) -> List[Tuple[float, float, float, float]]:
    cmap = plt.get_cmap("Set2")
    if hasattr(cmap, "colors") and len(cmap.colors) >= n:
        return list(cmap.colors[:n])
    fallback = plt.get_cmap("tab20c")
    return [fallback(v) for v in np.linspace(0, 1, n)]


def _configure_plot_style() -> None:
    plt.rcParams.update(
        {
            "axes.titlesize": 12,
            "axes.labelsize": 11,
            "legend.fontsize": 9,
            "font.size": 10,
        }
    )


def plot_farm_domain_shift(
    out_path: Path,
    farm_label: str,
    wind_hub: np.ndarray,
    power: np.ndarray,
    extreme_mask: np.ndarray,
    rng: np.random.Generator,
    max_samples: int = 8000,
) -> None:
    features = [
        ("wind_hub", wind_hub, "Wind speed at hub height (m/s)"),
        ("power", power, "Power (MW)"),
    ]

    colors = _get_palette(2)
    color_src, color_tgt = colors[0], colors[1]
    extreme_mask = np.asarray(extreme_mask, dtype=bool)

    fig, axes = plt.subplots(nrows=len(features), ncols=1, figsize=(9.5, 7.2))
    if len(features) == 1:
        axes = [axes]

    for ax, (_, values, xlabel) in zip(axes, features):
        vals = np.asarray(values, dtype=float)
        mask_src = (~extreme_mask) & np.isfinite(vals)
        mask_tgt = extreme_mask & np.isfinite(vals)
        src = vals[mask_src]
        tgt = vals[mask_tgt]

        src = _maybe_subsample(src, max_samples=max_samples, rng=rng)
        tgt = _maybe_subsample(tgt, max_samples=max_samples, rng=rng)

        if len(src) < 2 or len(tgt) < 2:
            ax.text(
                0.5,
                0.5,
                f"Not enough samples for KDE\nsource={len(src)}, target={len(tgt)}",
                ha="center",
                va="center",
            )
            ax.set_axis_off()
            continue

        lo = float(min(np.min(src), np.min(tgt)))
        hi = float(max(np.max(src), np.max(tgt)))
        if math.isclose(lo, hi):
            hi = lo + 1.0
        pad = 0.05 * (hi - lo)
        xs = np.linspace(lo - pad, hi + pad, 420)

        kde_src = gaussian_kde(src)
        kde_tgt = gaussian_kde(tgt)
        y_src = kde_src(xs)
        y_tgt = kde_tgt(xs)

        ax.plot(xs, y_src, color=color_src, lw=2.2, label="Source (normal)")
        ax.fill_between(xs, y_src, color=color_src, alpha=0.25)
        ax.plot(xs, y_tgt, color=color_tgt, lw=2.2, label="Target (extreme)")
        ax.fill_between(xs, y_tgt, color=color_tgt, alpha=0.25)

        ax.set_xlabel(xlabel)
        ax.set_ylabel("Density")
        ax.grid(True, alpha=0.25, linestyle="--")
        ax.legend(loc="upper right", frameon=False)

    fig.suptitle(f"Distribution shift: {farm_label} (Normal vs Extreme)", fontsize=13)
    fig.tight_layout(rect=[0.0, 0.0, 1.0, 0.95])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=300)
    plt.close(fig)


# -----------------------------
# Main
# -----------------------------


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--max-farms", type=int, default=6, help="Number of farms to process (capacity desc)")
    parser.add_argument("--farm-idx", type=int, default=None, help="Farm index for time-series output")
    parser.add_argument("--with-timeseries", action="store_true", help="Generate extreme-event time-series snippets")
    parser.add_argument("--feature-set", type=str, default="full", choices=["full", "wind_only", "simple"])
    parser.add_argument("--seed", type=int, default=0)
    args = parser.parse_args()

    _configure_plot_style()

    exp_dir = Path(__file__).resolve().parent
    out_dir = exp_dir / "outputs"
    plot_dir = out_dir / "plots"
    out_dir.mkdir(parents=True, exist_ok=True)
    plot_dir.mkdir(parents=True, exist_ok=True)

    data_dir = (exp_dir.parent / "data" / "wind_farm_data" / "data_processed" / "wind_farms").resolve()
    farms = list_wind_farms(data_dir)
    if not farms:
        raise RuntimeError(f"No .xlsx wind farm files found under {data_dir}")
    if args.max_farms <= 0:
        raise ValueError("--max-farms must be positive")
    if len(farms) < args.max_farms:
        raise RuntimeError(f"Need at least {args.max_farms} farms, found {len(farms)}")
    farms = farms[: args.max_farms]

    rng = np.random.default_rng(args.seed)

    if args.with_timeseries:
        if args.farm_idx is None:
            raise ValueError("--with-timeseries requires --farm-idx")
        if args.farm_idx < 0 or args.farm_idx >= len(farms):
            raise ValueError(f"--farm-idx out of range. Available: 0..{len(farms)-1}")
    for i, farm in enumerate(farms):
        print(f"[Load] Farm {i} | {farm.name} | capacity={farm.capacity_mw} MW")
        series, col_names = load_farm_timeseries_xlsx(farm.path, feature_set=args.feature_set)

        time_str = series.get("time", np.array([], dtype=object))
        power = series["power"]
        wind_hub = series["wind_hub"]
        temperature = series.get("temperature", None)

        masks = identify_extreme_events(
            power_mw=power,
            wind_hub_ms=wind_hub,
            temperature_c=temperature,
            nominal_capacity_mw=farm.capacity_mw,
        )

        meta = {
            "farm_idx": i,
            "farm_name": farm.name,
            "farm_file": str(farm.path),
            "nominal_capacity_mw": farm.capacity_mw,
            "feature_set": args.feature_set,
            "columns": col_names,
            "n_samples": int(len(power)),
            "n_extreme": int(np.sum(masks.extreme)),
            "n_ramp": int(np.sum(masks.ramp)),
            "n_cutout": int(np.sum(masks.cutout)),
            "n_stat": int(np.sum(masks.stat)),
            "n_temp": int(np.sum(masks.temp)),
            "q95_wind_hub_ms": masks.q95_wind,
            "q95_power_mw": masks.q95_power,
            "ramp_threshold_mw": masks.ramp_threshold_mw,
            "cutout_threshold_ms": 25.0,
            "temp_threshold_c": {"low": -10.0, "high": 35.0},
            "seed": args.seed,
        }
        (out_dir / f"farm{i}_meta.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")

        csv_path = out_dir / f"farm{i}_events.csv"
        with csv_path.open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["index", "time", "wind_hub_ms", "power_mw", "is_extreme", "is_ramp", "is_cutout"])
            for j in range(len(power)):
                w.writerow(
                    [
                        j,
                        time_str[j] if len(time_str) else "",
                        float(wind_hub[j]) if np.isfinite(wind_hub[j]) else "",
                        float(power[j]) if np.isfinite(power[j]) else "",
                        int(bool(masks.extreme[j])),
                        int(bool(masks.ramp[j])),
                        int(bool(masks.cutout[j])),
                    ]
                )

        print(
            f"[Screening] farm={i} total={len(power)} | extreme={np.sum(masks.extreme)} "
            f"(ramp={np.sum(masks.ramp)}, cutout={np.sum(masks.cutout)})"
        )

        plot_farm_domain_shift(
            plot_dir / f"farm{i}_domain_shift.png",
            farm_label=f"Farm {i + 1} ({farm.capacity_mw} MW)",
            wind_hub=wind_hub,
            power=power,
            extreme_mask=masks.extreme,
            rng=rng,
        )

        if args.with_timeseries and i == args.farm_idx:
            pre, post = 48, 48
            cut_idx = np.flatnonzero(masks.cutout)
            ramp_idx = np.flatnonzero(masks.ramp)

            cut_segs = _select_non_overlapping_segments(cut_idx, n=3, pre=pre, post=post, total_len=len(power))
            ramp_segs = _select_non_overlapping_segments(ramp_idx, n=3, pre=pre, post=post, total_len=len(power))

            if not cut_segs:
                print("[TimeSeries] No cut-out segments found (wind_hub > 25m/s).")
            for k, (ei, s, e) in enumerate(cut_segs, start=1):
                plot_timeseries_segment(
                    plot_dir / f"farm{i}_cutout_seg{k}.png",
                    time_str=time_str,
                    wind_hub=wind_hub,
                    power=power,
                    event_i=ei,
                    start=s,
                    end=e,
                    title=f"Cut-out event snippet (wind_hub > 25 m/s) | {farm.capacity_mw}MW",
                )

            if not ramp_segs:
                print("[TimeSeries] No ramp segments found (|ΔP| > 5% capacity per 15min).")
            for k, (ei, s, e) in enumerate(ramp_segs, start=1):
                plot_timeseries_segment(
                    plot_dir / f"farm{i}_ramp_seg{k}.png",
                    time_str=time_str,
                    wind_hub=wind_hub,
                    power=power,
                    event_i=ei,
                    start=s,
                    end=e,
                    title=f"Ramp event snippet (|ΔP| > {masks.ramp_threshold_mw:.2f} MW) | {farm.capacity_mw}MW",
                )

    print(f"[Done] Outputs saved under: {out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
