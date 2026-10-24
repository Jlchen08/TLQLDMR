#!/usr/bin/env python3
"""Read-only verification of one corrected-validation run; never fits a model.

Recomputes predictions from saved models and raw inputs, metrics, training-only
preprocessing, and calibration-only endpoints independently of the runner.
Prints a JSON audit. Exit 0 means passed (or an allowed, explicitly partial
audit); exit 1 means a mismatch. Random overlapping windows remain a limitation.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.spatial.distance import cdist
from threadpoolctl import threadpool_limits

ROOT = Path(__file__).resolve().parents[1]


def sha256(path):
    result = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            result.update(block)
    return result.hexdigest()


def read_json(path):
    return json.loads(path.read_text())


def read_csv(path):
    with path.open(newline="") as stream:
        rows = list(csv.DictReader(stream))
    if not rows:
        raise ValueError(f"Empty CSV: {path}")
    return {key: np.asarray([row[key] for row in rows], dtype=(str if key == "time" else float))
            for key in rows[0]}


def load_raw(path, feature_names):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        iterator = workbook.worksheets[0].iter_rows(values_only=True)
        names = [str(value) for value in next(iterator)]
        rows = list(iterator)
    finally:
        workbook.close()
    # Match pandas' exclusion of entirely blank trailing spreadsheet rows.
    while rows and all(value is None for value in rows[-1]):
        rows.pop()

    def numeric(name):
        index = names.index(name)
        return np.array([float(row[index]) if row[index] is not None else np.nan
                         for row in rows])

    columns = []
    for name in feature_names:
        if name.startswith("sin(") or name.startswith("cos("):
            values = np.deg2rad(numeric(name[4:-1]))
            columns.append(np.sin(values) if name.startswith("sin(") else np.cos(values))
        else:
            columns.append(numeric(name))
    power = next(name for name in names if "power" in name.lower() and "mw" in name.lower())
    time_index = next((index for index, name in enumerate(names) if "time" in name.lower()), None)
    timestamps = (np.asarray([str(row[time_index]) for row in rows]) if time_index is not None
                  else np.arange(len(rows)).astype(str))
    return np.column_stack(columns), numeric(power), timestamps


def point_metrics(y, prediction):
    errors = y - prediction
    return {"rmse": float(np.sqrt(np.mean(errors**2))),
            "mae": float(np.mean(np.abs(errors))),
            "r2": float(1 - np.sum(errors**2) / np.sum((y - y.mean())**2))}


def interval_metrics(y, lower, upper, alpha):
    width = upper - lower
    coverage = float(np.mean((lower <= y) & (y <= upper)))
    mean_width = float(np.mean(width))
    normalized = mean_width / float(np.max(y) - np.min(y))
    penalty = 1 if coverage >= 1 - alpha else 1 + 0.3 * math.exp(-2 * (coverage - (1 - alpha)))
    score = width + (2 / alpha) * (np.maximum(lower - y, 0) + np.maximum(y - upper, 0))
    return {"picp": coverage, "mpiw": mean_width, "pinaw": normalized,
            "cwc": normalized * penalty, "winkler": float(np.mean(score))}


class Audit:
    def __init__(self):
        self.errors = []
        self.checks = 0

    def require(self, condition, message):
        self.checks += 1
        if not condition:
            self.errors.append(message)

    def close(self, actual, expected, message, atol=1e-8, rtol=1e-9):
        actual, expected = np.asarray(actual), np.asarray(expected)
        equal = (actual.shape == expected.shape and np.isfinite(actual).all()
                 and np.isfinite(expected).all()
                 and np.allclose(actual, expected, atol=atol, rtol=rtol))
        self.require(equal, message)

    def ids(self, values, message):
        values = np.asarray(values)
        valid = (values.ndim == 1 and np.isfinite(values).all()
                 and np.all(values >= 0) and np.all(values == values.astype(np.int64))
                 and len(np.unique(values)) == len(values))
        self.require(valid, message)
        if not valid:
            raise ValueError(message)
        return values.astype(np.int64)


def verify_preprocessing(audit, prefix, saved, raw_x, raw_y, window, source, target):
    ids = np.r_[source, target]
    counts = np.zeros(len(raw_x) + 1, dtype=np.int64)
    np.add.at(counts, ids, 1)
    np.add.at(counts, ids + window, -1)
    used = np.cumsum(counts[:-1]) > 0
    selected = np.where(np.isfinite(raw_x[used]), raw_x[used], np.nan)
    finite_counts = np.isfinite(selected).sum(axis=0)
    imputation = np.divide(np.nansum(selected, axis=0), finite_counts,
                           out=np.zeros(raw_x.shape[1]), where=finite_counts > 0)
    filled = np.where(np.isfinite(raw_x), raw_x, imputation)
    mean = filled[used].mean(axis=0)
    scale = filled[used].std(axis=0)
    scale[scale == 0] = 1
    training_labels = raw_y[ids + window]
    label_scale = float(training_labels.std()) or 1.0
    for key, expected in (("feature_imputation_mean", imputation), ("feature_mean", mean),
                          ("feature_scale", scale), ("label_mean", training_labels.mean()),
                          ("label_scale", label_scale)):
        audit.close(saved[key], expected, f"{prefix}: {key} is not from fitting samples only")
    audit.require(saved["fit_label_count"] == len(ids), f"{prefix}: fit label count differs")
    audit.require(saved["unique_training_input_rows"] == int(used.sum()),
                  f"{prefix}: training input-row count differs")
    # Use the saved transform when reconstructing the saved model's predictions.
    filled = np.where(np.isfinite(raw_x), raw_x, saved["feature_imputation_mean"])
    return ((filled - saved["feature_mean"]) / saved["feature_scale"]).astype(np.float32)


def model_prediction(path, scaled_x, ids, window, options, preprocessing, audit, prefix):
    with np.load(path, allow_pickle=False) as model:
        w, b = model["w"], model["b"]
        landmarks, normalization = model["landmarks"], model["normalization"]
    m = len(landmarks)
    audit.require(w.shape == (m,) and normalization.shape == (m, m),
                  f"{prefix}: incompatible saved parameter shapes")
    audit.require(all(np.isfinite(a).all() for a in (w, b, landmarks, normalization)),
                  f"{prefix}: nonfinite saved parameters")
    inputs = scaled_x[ids[:, None] + np.arange(window)[None, :]].reshape(len(ids), -1)
    audit.require(inputs.shape[1] == landmarks.shape[1], f"{prefix}: feature dimension differs")
    beta = normalization @ w
    predictions = []
    for offset in range(0, len(inputs), 256):
        kernel = np.exp(-options["kernel_gamma"] * cdist(
            inputs[offset:offset + 256], landmarks, metric="sqeuclidean"))
        predictions.append(kernel @ beta + float(b))
    scaled = np.concatenate(predictions)
    return scaled * preprocessing["label_scale"] + preprocessing["label_mean"], m


def expected_options(config, task, variant, tau):
    params = config["params"]
    options = {name: params[name] for name in ("lambda1", "lambda2", "C_S", "C_T", "kernel_gamma")}
    options.update({"tau": tau, "variance_mode": "centered_residual", "asym_scale": 0.0,
                    "n_components": int(params["nystrom_n_components"]),
                    "learning_rate": params["nystrom_lr"], "epochs": int(params["nystrom_epochs"]),
                    "batch_size": int(params["nystrom_batch_size"])})
    if variant == "NoVariance":
        options["lambda1"] = 0.0
    if variant in ("NoMMD", "TargetOnly"):
        options["lambda2"] = 0.0
    if variant == "TargetOnly":
        options["C_S"] = 0.0
    return options


def verify_branch(audit, out, task, variant, manifest, split, raw_x, raw_y):
    prefix = out.name
    config = manifest["configs"][task]
    window = int(config["window_size"])
    alpha = manifest["protocol"]["alpha"]
    source, train, test = split
    active_source = source if variant != "TargetOnly" else np.array([], dtype=np.int64)
    n_cal = max(1, int(len(train) * manifest["protocol"]["calibration_ratio"])) if task == "interval" else 0
    fit, cal = (train[:-n_cal], train[-n_cal:]) if n_cal else (train, np.array([], dtype=np.int64))
    with np.load(out / "fit_ids.npz", allow_pickle=False) as saved_ids:
        for key, expected in (("source_window_ids", active_source), ("target_fit_window_ids", fit),
                              ("calibration_window_ids", cal)):
            actual = audit.ids(saved_ids[key], f"{prefix}: invalid {key}")
            audit.require(np.array_equal(actual, expected), f"{prefix}: {key} differs from declared split")
    for first, second, names in ((active_source, fit, "source/fit"), (active_source, cal, "source/cal"),
                                 (active_source, test, "source/test"), (fit, cal, "fit/cal"),
                                 (fit, test, "fit/test"), (cal, test, "cal/test")):
        audit.require(not len(np.intersect1d(first, second)), f"{prefix}: {names} window IDs overlap")

    result = read_json(out / "result.json")
    audit.require(result["task"] == task and result["variant"] == variant,
                  f"{prefix}: result identity differs")
    prep = read_json(out / "preprocessing.json")
    scaled_x = verify_preprocessing(audit, prefix, prep, raw_x, raw_y, window, active_source, fit)
    table = read_csv(out / "test_predictions.csv")
    audit.require(np.array_equal(table["window_id"], test), f"{prefix}: test row IDs/order differ")
    audit.require(np.array_equal(table["raw_target_row"], test + window), f"{prefix}: raw test rows differ")
    audit.close(table["y_true"], raw_y[test + window], f"{prefix}: test labels differ from raw file")
    cal_table = read_csv(out / "calibration.csv") if n_cal else None
    if cal_table is not None:
        audit.require(np.array_equal(cal_table["window_id"], cal), f"{prefix}: calibration IDs/order differ")
        audit.require(np.array_equal(cal_table["raw_target_row"], cal + window), f"{prefix}: raw calibration rows differ")
        audit.close(cal_table["y_true"], raw_y[cal + window], f"{prefix}: calibration labels differ from raw file")
    taus = [0.5] if task == "point" else [float(config.get("tau_low", 0.2)), float(config.get("tau_high", 0.8))]
    audit.require(len(result["fits"]) == len(taus), f"{prefix}: model count differs")
    for index, tau in enumerate(taus):
        options = expected_options(config, task, variant, tau)
        fit_record = result["fits"][index]
        audit.require(fit_record["options"] == options, f"{prefix}: model options changed from fixed configuration")
        prediction, m = model_prediction(out / f"model_tau_{tau}.npz", scaled_x, np.r_[test, cal],
                                          window, options, prep, audit, prefix)
        audit.require(fit_record["effective_landmarks"] == m, f"{prefix}: landmark count differs")
        pool_size = len(active_source) if len(active_source) else len(fit)
        audit.require(m == min(2000, pool_size, options["n_components"]), f"{prefix}: unexpected landmark count")
        column = "prediction" if task == "point" else ("raw_tau_low" if index == 0 else "raw_tau_high")
        audit.close(table[column], prediction[:len(test)], f"{prefix}: {column} differs from saved model (possible postprocessing)", atol=2e-6)
        if n_cal:
            audit.close(cal_table[column], prediction[len(test):], f"{prefix}: calibration {column} differs from saved model", atol=2e-6)

    if task == "point":
        audit.require(result["calibration"] is None, f"{prefix}: unexpected point calibration")
        metrics = point_metrics(table["y_true"], table["prediction"])
    else:
        for name, data in (("test", table), ("calibration", cal_table)):
            low, high = np.minimum(data["raw_tau_low"], data["raw_tau_high"]), np.maximum(data["raw_tau_low"], data["raw_tau_high"])
            audit.close(data["ordered_lower"], low, f"{prefix}: {name} lower ordering changed")
            audit.close(data["ordered_upper"], high, f"{prefix}: {name} upper ordering changed")
        scores = np.maximum.reduce((cal_table["ordered_lower"] - cal_table["y_true"],
                                     cal_table["y_true"] - cal_table["ordered_upper"], np.zeros(n_cal)))
        rank = math.ceil((n_cal + 1) * (1 - alpha))
        if rank > n_cal:
            raise ValueError(f"{prefix}: insufficient calibration observations for finite-sample rank")
        qhat = float(np.sort(scores)[rank - 1])
        audit.close(cal_table["score"], scores, f"{prefix}: calibration scores differ")
        calibration = result["calibration"]
        audit.require(calibration["rank"] == rank and calibration["n_cal"] == n_cal,
                      f"{prefix}: finite-sample calibration rank/count differ")
        audit.close(calibration["q_hat"], qhat, f"{prefix}: q_hat is not the calibration-only order statistic")
        audit.require(calibration["q_scale"] == 1, f"{prefix}: undeclared q_scale")
        audit.close(table["lower"], table["ordered_lower"] - qhat, f"{prefix}: lower endpoint has extra postprocessing")
        audit.close(table["upper"], table["ordered_upper"] + qhat, f"{prefix}: upper endpoint has extra postprocessing")
        raw_metrics = interval_metrics(table["y_true"], table["ordered_lower"], table["ordered_upper"], alpha)
        for name, value in raw_metrics.items():
            audit.close(calibration["raw_metrics"][name], value, f"{prefix}: raw {name} differs")
        metrics = interval_metrics(table["y_true"], table["lower"], table["upper"], alpha)
    for name, value in metrics.items():
        audit.close(result["metrics"][name], value, f"{prefix}: reported {name} differs from per-sample recomputation")
    return metrics


def verify(run, allow_partial=False):
    audit = Audit()
    manifest = read_json(run / "manifest.json")
    status = manifest.get("status")
    audit.require(status == "complete" or (allow_partial and status == "running"),
                  f"Manifest status is {status!r}, not complete")
    if status == "complete":
        audit.require(bool(manifest.get("finished_utc")), "Complete manifest lacks finished_utc")
    hashes = {}
    for label, path, expected in [("data", Path(manifest["data_file"]), manifest["data_sha256"])] + [
            (name, ROOT / name, value) for name, value in manifest["source_sha256"].items()]:
        actual = sha256(path) if path.is_file() else None
        hashes[label] = {"expected": expected, "actual": actual, "matches": actual == expected}
        audit.require(actual == expected, f"SHA256 mismatch or missing file: {label}")
    raw_x, raw_y, timestamps = load_raw(Path(manifest["data_file"]), manifest["feature_names"])
    config = manifest["configs"][manifest["tasks"][0]]
    window = int(config["window_size"])
    with np.load(run / "split_ids.npz", allow_pickle=False) as saved:
        source = audit.ids(saved["source_window_ids"], "Invalid source split IDs")
        train = audit.ids(saved["target_training_window_ids"], "Invalid training split IDs")
        test = audit.ids(saved["test_window_ids"], "Invalid test split IDs")
        audit.require(np.array_equal(saved["test_raw_target_rows"], test + window), "Split raw target rows differ")
        # The runner saved timestamps as an object array. Reconstruct from the
        # raw workbook instead of enabling pickle loading for a data artifact.
        test_times = timestamps[test + window]
    all_ids = np.r_[source, train, test]
    audit.require(np.array_equal(np.sort(all_ids), np.arange(len(raw_y) - window)),
                  "Source/training/test IDs are not a disjoint complete window partition")
    target = np.sort(np.r_[train, test])
    shuffled = target[np.random.default_rng(int(config.get("split_seed", 0))).permutation(len(target))]
    split_at = int(len(target) * config["target_train_ratio"])
    audit.require(np.array_equal(train, shuffled[:split_at]) and np.array_equal(test, shuffled[split_at:]),
                  "Target split differs from declared seed/permutation/ratio")
    for key, count in (("source", len(source)), ("target_train_pool", len(train)), ("test", len(test))):
        audit.require(manifest["counts"][key] == count, f"Manifest count differs: {key}")

    completed, pending, metrics = [], [], {}
    for task in manifest["tasks"]:
        for variant in manifest["variants"]:
            name = f"{task}_{variant}"
            out = run / name
            if not (out / "result.json").is_file():
                pending.append(name)
                audit.require(allow_partial and status == "running", f"Missing completed branch: {name}")
                continue
            before = len(audit.errors)
            try:
                metrics[name] = verify_branch(audit, out, task, variant, manifest, (source, train, test), raw_x, raw_y)
                table = read_csv(out / "test_predictions.csv")
                audit.require(np.array_equal(table["time"], test_times), f"{name}: test timestamps/order differ")
                completed.append({"branch": name, "passed": len(audit.errors) == before})
            except (KeyError, ValueError, OSError, IndexError) as error:
                audit.errors.append(f"{name}: {type(error).__name__}: {error}")
                completed.append({"branch": name, "passed": False})
    if metrics:
        with (run / "summary_metrics.csv").open(newline="") as stream:
            summary = list(csv.DictReader(stream))
        seen = set()
        for row in summary:
            name = f"{row['task']}_{row['variant']}"
            # A live runner may finish a previously pending branch while the
            # verifier recomputes other models. Leave that new row for the
            # next audit instead of misreporting concurrent progress as error.
            if allow_partial and status == "running" and name in pending:
                continue
            audit.require(name not in seen, f"Duplicate summary row: {name}")
            seen.add(name)
            audit.require(name in metrics, f"Summary row lacks a verified branch: {name}")
            for key, value in metrics.get(name, {}).items():
                audit.close(float(row[key]), value, f"{name}: summary {key} differs")
        audit.require(seen == set(metrics), "Summary rows do not match completed branch results")
    return {"status": "failed" if audit.errors else ("passed" if status == "complete" else "partial"),
            "run_directory": str(run.resolve()), "manifest_status": status, "checks": audit.checks,
            "hashes": hashes, "completed_branches": completed, "pending_branches": pending,
            "recomputed_metrics": metrics, "errors": audit.errors,
            "scope": "Read-only arithmetic/provenance/model-output audit; no training or parameter selection.",
            "limitations": ["Disjoint window IDs do not imply disjoint overlapping input rows or chronological independence.",
                            "Artifact consistency cannot establish untouched historical hyperparameter selection or prove past execution history.",
                            "Hash checks compare current source/data to the run manifest; earlier modified source is reported as drift."]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("run_directory", type=Path)
    parser.add_argument("--allow-partial", action="store_true", help="Audit finished branches while a manifest is running.")
    args = parser.parse_args()
    try:
        with threadpool_limits(limits=1):
            report = verify(args.run_directory, args.allow_partial)
    except (OSError, KeyError, ValueError, IndexError) as error:
        report = {"status": "failed", "errors": [f"{type(error).__name__}: {error}"]}
    print(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False))
    raise SystemExit(1 if report["status"] == "failed" else 0)


if __name__ == "__main__":
    main()
